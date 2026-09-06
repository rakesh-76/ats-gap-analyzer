"""Unit tests for pipeline.extract (Stage 2: text extraction & cleaning)."""

from __future__ import annotations

import io

import docx
import pytest

from ats_gap_analyzer.pipeline.extract import (
    Bullet,
    ExtractedDocument,
    ExtractionError,
    extract_job_description,
    extract_resume_text,
)


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    """Build an in-memory .docx file to use as a test fixture."""
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _texts(bullets: list[Bullet]) -> list[str]:
    return [b.text for b in bullets]


class TestExtractResumeText:
    def test_unsupported_file_type_raises(self) -> None:
        with pytest.raises(ExtractionError):
            extract_resume_text(b"irrelevant", "txt")

    def test_corrupt_pdf_raises_extraction_error(self) -> None:
        with pytest.raises(ExtractionError):
            extract_resume_text(b"not a real pdf", "pdf")

    def test_corrupt_docx_raises_extraction_error(self) -> None:
        with pytest.raises(ExtractionError):
            extract_resume_text(b"not a real docx", "docx")

    def test_docx_extraction_returns_cleaned_bullets(self) -> None:
        file_bytes = _make_docx_bytes(
            ["Jane Doe", "", "Experience", "Built things at Acme Corp", "Page 1"]
        )

        result = extract_resume_text(file_bytes, ".DOCX")  # dot + casing tolerated

        assert isinstance(result, ExtractedDocument)
        bullet_texts = _texts(result.bullets)
        assert "Jane Doe" in bullet_texts
        assert "Built things at Acme Corp" in bullet_texts
        assert "" not in bullet_texts  # blank lines stripped
        assert "Page 1" not in bullet_texts  # page-number boilerplate stripped

    def test_docx_sections_group_bullets_under_headers(self) -> None:
        file_bytes = _make_docx_bytes(
            ["Summary", "Enthusiastic engineer", "Skills", "\u2022 Python", "\u2022 SQL"]
        )

        result = extract_resume_text(file_bytes, "docx")

        assert _texts(result.sections["summary"]) == ["Enthusiastic engineer"]
        assert _texts(result.sections["skills"]) == ["Python", "SQL"]
        assert "content" not in result.sections  # nothing appeared before a header

    def test_docx_table_text_is_included_as_separate_cells(self) -> None:
        document = docx.Document()
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Python"
        table.rows[0].cells[1].text = "5 years"
        buffer = io.BytesIO()
        document.save(buffer)

        result = extract_resume_text(buffer.getvalue(), "docx")

        # Table cells are independent data, not a wrapped sentence, so
        # each must survive the bullet-merge step as its own item.
        bullet_texts = _texts(result.bullets)
        assert "Python" in bullet_texts
        assert "5 years" in bullet_texts

    def test_no_recognized_headers_merges_unmarked_wrapped_lines(self) -> None:
        file_bytes = _make_docx_bytes(["Just some text", "with no headers at all"])

        result = extract_resume_text(file_bytes, "docx")

        assert _texts(result.sections["content"]) == [
            "Just some text with no headers at all"
        ]

    def test_wrapped_bullet_is_merged_into_one_item(self) -> None:
        """Regression test: a bullet that wraps onto a second line must not
        become two separate, incomplete fragments."""
        file_bytes = _make_docx_bytes(
            [
                "Experience",
                "\u2022 Collaborated with the team to design and implement systems for",
                "efficient data storage and retrieval.",
            ]
        )

        result = extract_resume_text(file_bytes, "docx")

        assert _texts(result.sections["experience"]) == [
            "Collaborated with the team to design and implement systems for "
            "efficient data storage and retrieval."
        ]

    def test_header_after_unterminated_line_still_starts_fresh(self) -> None:
        """Regression test: a section header immediately after a
        contact-info line with no ending punctuation must still be
        recognized as its own header, not swallowed into that line."""
        file_bytes = _make_docx_bytes(
            [
                "Rakesh Ravuri",
                "linkedin.com/in/example",
                "Professional Summary",
                "Built things.",
            ]
        )

        result = extract_resume_text(file_bytes, "docx")

        assert _texts(result.sections["summary"]) == ["Built things."]
        assert _texts(result.sections["content"]) == ["Rakesh Ravuri linkedin.com/in/example"]

    def test_certifications_and_achievements_header_is_recognized(self) -> None:
        file_bytes = _make_docx_bytes(
            [
                "Professional Summary",
                "Great engineer.",
                "Certifications & Achievements",
                "AWS Certified.",
            ]
        )

        result = extract_resume_text(file_bytes, "docx")

        assert _texts(result.sections["summary"]) == ["Great engineer."]
        assert _texts(result.sections["certifications"]) == ["AWS Certified."]

    def test_cid_font_artifacts_are_stripped(self) -> None:
        """Regression test: icon-font glyphs with no Unicode mapping leak
        into extracted text as "(cid:NNN)" tokens and must be removed."""
        file_bytes = _make_docx_bytes(
            ["Contact: (cid:239) linkedin.com/in/example (cid:131) 555-1234"]
        )

        result = extract_resume_text(file_bytes, "docx")

        assert not any("cid:" in b.text for b in result.bullets)
        assert result.full_text == "Contact: linkedin.com/in/example 555-1234"

    def test_contact_block_is_flagged_as_heading(self) -> None:
        file_bytes = _make_docx_bytes(
            [
                "Jane Doe",
                "linkedin.com/in/jane 555-1234",
                "Experience",
                "\u2022 Did things that mattered.",
            ]
        )

        result = extract_resume_text(file_bytes, "docx")

        assert result.sections["content"][0].is_heading is True

    def test_job_title_and_company_line_is_flagged_as_heading(self) -> None:
        file_bytes = _make_docx_bytes(
            [
                "Experience",
                "Software Engineer, Acme Corp",
                "\u2022 Shipped a real feature that mattered.",
            ]
        )

        result = extract_resume_text(file_bytes, "docx")

        experience = result.sections["experience"]
        assert experience[0] == Bullet(text="Software Engineer, Acme Corp", is_heading=True)
        assert experience[1].is_heading is False

    def test_second_project_title_mid_section_is_flagged_as_heading(self) -> None:
        """Regression test: a second (or later) entry title within the same
        multi-entry section — not just the one right after the section
        header — must still be recognized as a heading."""
        file_bytes = _make_docx_bytes(
            [
                "Projects",
                "First Project | Python",
                "\u2022 Did the first thing.",
                "Second Project | SQL",
                "\u2022 Did the second thing.",
            ]
        )

        result = extract_resume_text(file_bytes, "docx")

        projects = result.sections["projects"]
        assert [b.is_heading for b in projects] == [True, False, True, False]

    def test_prose_summary_is_not_flagged_as_heading(self) -> None:
        """Regression test for a real bug: a resume where Experience uses
        bullet markers must NOT cause the unmarked Summary paragraph
        (which never uses markers, by normal convention) to be
        misclassified as a heading."""
        file_bytes = _make_docx_bytes(
            [
                "Professional Summary",
                "Experienced engineer with a track record of shipping things.",
                "Experience",
                "Software Engineer, Acme Corp",
                "\u2022 Shipped a real feature that mattered.",
            ]
        )

        result = extract_resume_text(file_bytes, "docx")

        assert result.sections["summary"][0].is_heading is False
        assert result.sections["experience"][0].is_heading is True  # job title still flagged

    def test_education_single_line_is_not_flagged_as_heading(self) -> None:
        file_bytes = _make_docx_bytes(
            ["Education", "State University, B.S. Computer Science, 2020"]
        )

        result = extract_resume_text(file_bytes, "docx")

        assert result.sections["education"][0].is_heading is False

    def test_resume_with_no_markers_anywhere_falls_back_to_colon_only(self) -> None:
        """Known-limitation regression guard: if a resume never uses bullet
        markers at all, the marker signal is meaningless, so nothing gets
        flagged as a heading purely for lacking one."""
        file_bytes = _make_docx_bytes(
            [
                "Experience",
                "Software Engineer, Acme Corp",
                "Shipped a real feature that mattered.",
            ]
        )

        result = extract_resume_text(file_bytes, "docx")

        assert all(not b.is_heading for b in result.sections["experience"])


class TestExtractJobDescription:
    def test_cleans_and_splits_pasted_text(self) -> None:
        raw = "We are hiring a senior engineer.\n\n  Requires Python and SQL.  \nPage 2"

        result = extract_job_description(raw)

        assert _texts(result.bullets) == [
            "We are hiring a senior engineer.",
            "Requires Python and SQL.",
        ]
        assert result.sections == {}  # sections are never computed for JDs

    def test_empty_text_returns_empty_document(self) -> None:
        result = extract_job_description("")

        assert result.bullets == []
        assert result.full_text == ""

    def test_internal_label_is_flagged_as_heading_but_requirements_are_not(self) -> None:
        raw = (
            "What You'll Do:\n"
            "Build things that matter.\n"
            "Skills: Python, SQL, and a willingness to learn."
        )

        result = extract_job_description(raw)

        assert result.bullets[0] == Bullet(text="What You'll Do:", is_heading=True)
        assert result.bullets[1].is_heading is False
        # "Skills:" mid-sentence must not trigger heading — only a bullet
        # that *ends* in a bare colon should.
        assert result.bullets[2].is_heading is False
