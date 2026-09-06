"""Stage 2 of the ATS Gap Analyzer pipeline: text extraction & cleaning.

Converts raw resume bytes (PDF or DOCX) and a raw job description string
into a clean, structured representation that later pipeline stages can
work with. Contains no Streamlit-specific code so it can be developed,
reasoned about, and tested independently of the UI layer.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

__all__ = [
    "Bullet",
    "ExtractedDocument",
    "ExtractionError",
    "extract_resume_text",
    "extract_job_description",
]


class ExtractionError(Exception):
    """Raised when a resume file cannot be parsed into text."""


@dataclass(frozen=True)
class Bullet:
    """A single cleaned, reflowed line of resume or job description text.

    Attributes:
        text: The cleaned text of this bullet.
        is_heading: True if this looks like a title/label line (a job
            title + company line, a project or certification title, a
            resume's contact-info block, or a short internal JD label
            like "What You'll Do:") rather than a genuine achievement or
            requirement statement. Headings still appear in `bullets`
            and `sections` — project titles in particular often list
            real technologies worth keeping for keyword extraction — but
            later stages doing sentence-level semantic matching should
            generally skip them.
    """

    text: str
    is_heading: bool = False


@dataclass
class ExtractedDocument:
    """Structured, cleaned text produced by this stage.

    Attributes:
        full_text: The complete cleaned text, bullets joined by newlines.
        bullets: Individual bullets, cleaned, with wrapped lines merged
            back into whole sentences/titles, and empty, page-number, or
            pure section-header lines removed.
        sections: Best-effort mapping of section name -> bullets
            belonging to that section. Bullets before any recognized
            header, or all bullets when no headers are found at all,
            are grouped under "content". Always empty for job
            descriptions, since section detection is resume-specific.
    """

    full_text: str
    bullets: list[Bullet] = field(default_factory=list)
    sections: dict[str, list[Bullet]] = field(default_factory=dict)


# Common resume section headers, mapped to a canonical name. Matching is
# done against the lowercased, punctuation-stripped bullet text, so
# real-world variants (e.g. "Professional Summary" vs. "Summary") need
# to be listed explicitly rather than relying on partial matching.
_SECTION_HEADERS: dict[str, tuple[str, ...]] = {
    "summary": ("summary", "professional summary", "objective", "profile"),
    "experience": (
        "experience",
        "work experience",
        "employment history",
        "professional experience",
    ),
    "education": ("education", "academic background"),
    "skills": ("skills", "technical skills", "core competencies"),
    "projects": ("projects", "personal projects"),
    "certifications": (
        "certifications",
        "licenses",
        "certifications & achievements",
        "certifications and achievements",
    ),
}

_HEADER_LOOKUP: dict[str, str] = {
    alias: canonical for canonical, aliases in _SECTION_HEADERS.items() for alias in aliases
}

# Sections where resumes conventionally alternate "entry title line" (job
# title + company, project name + tech stack, certification name) with one
# or more marker-prefixed detail bullets. An unmarked line is only treated
# as a title/heading within these sections (plus "content", which covers
# the contact-info block before any header appears) — Summary, Education,
# and Skills are typically flat prose or lists with no such title
# substructure, and an unmarked line there is genuine content, not a title.
_MULTI_ENTRY_SECTIONS = frozenset({"content", "experience", "projects", "certifications"})

# Matches lines that are only a page number, optionally in "Page X" or
# "X of Y" / "X/Y" form (case-insensitive) — common PDF/DOCX boilerplate.
_PAGE_NUMBER_RE = re.compile(r"^\s*(page\s+)?\d+(\s*(of|/)\s*\d+)?\s*$", re.IGNORECASE)

# Font-glyph artifacts PDF text extraction leaves behind when a font (often
# an icon font used for contact-info symbols) has no Unicode mapping for a
# character, e.g. "(cid:239)" in place of a LinkedIn icon.
_CID_ARTIFACT_RE = re.compile(r"\(cid:\d+\)")

# A line that starts with a bullet glyph or a numbered-list marker is
# always the start of a new item, regardless of surrounding punctuation.
_LEADING_BULLET_RE = re.compile(r"^(?:[•▪‣∙◦*\-]|\d+[.)])\s*")

# Punctuation that plausibly ends a complete sentence/item. A line ending
# in one of these is assumed to be "finished" — the next line starts a
# new bullet rather than continuing this one.
_TERMINAL_PUNCTUATION = ".!?:;"


def extract_resume_text(file_bytes: bytes, file_type: str) -> ExtractedDocument:
    """Extract and clean text from an uploaded resume file.

    Args:
        file_bytes: Raw file content, exactly as received from the upload
            widget.
        file_type: The resume format — "pdf" or "docx". Case-insensitive;
            a leading dot (e.g. ".pdf") is tolerated.

    Returns:
        An ExtractedDocument with full_text, bullets, and best-effort
        section labels.

    Raises:
        ExtractionError: If file_type is unsupported, or the file cannot
            be parsed (corrupt, empty, or unreadable content).
    """
    normalized_type = file_type.lower().lstrip(".")

    if normalized_type == "pdf":
        raw_lines = _extract_pdf_lines(file_bytes)
    elif normalized_type == "docx":
        raw_lines = _extract_docx_lines(file_bytes)
    else:
        raise ExtractionError(
            f"Unsupported resume file type: {file_type!r}. Expected 'pdf' or 'docx'."
        )

    cleaned = _clean_lines(raw_lines)
    bullets, sections = _process_resume_lines(cleaned)
    return ExtractedDocument(
        full_text="\n".join(b.text for b in bullets), bullets=bullets, sections=sections
    )


def extract_job_description(raw_text: str) -> ExtractedDocument:
    """Clean and structure a pasted job description.

    Job descriptions are provided as pasted text rather than an uploaded
    file (see the project spec), so this only needs to clean the text and
    reflow it into complete bullets — there is no file format to parse
    and no resume-style sections to detect.

    Args:
        raw_text: The raw, pasted job description text.

    Returns:
        An ExtractedDocument with full_text and bullets populated;
        sections is always empty.
    """
    cleaned = _clean_lines(raw_text.splitlines())
    bullets = _process_jd_lines(cleaned)
    return ExtractedDocument(full_text="\n".join(b.text for b in bullets), bullets=bullets, sections={})


def _extract_pdf_lines(file_bytes: bytes) -> list[str]:
    """Extract raw text lines from a PDF, page by page."""
    import pdfplumber  # imported here to keep module import light if unused

    lines: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                lines.extend(text.splitlines())
    except Exception as exc:
        # pdfplumber/pdfminer raise a variety of exception types for
        # malformed PDFs; converting all of them to ExtractionError keeps
        # this a clean, single failure mode for callers to handle.
        raise ExtractionError(f"Could not parse PDF resume: {exc}") from exc

    return lines


def _extract_docx_lines(file_bytes: bytes) -> list[str]:
    """Extract raw text lines from a DOCX, including any table cells."""
    import docx  # imported here to keep module import light if unused

    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ExtractionError(f"Could not parse DOCX resume: {exc}") from exc

    lines: list[str] = [paragraph.text for paragraph in document.paragraphs]

    # Some resumes use table layouts (e.g. for skills matrices). Table
    # cells are independent, discrete pieces of content — never a
    # sentence wrapped across lines — so each is tagged with a bullet
    # marker to force it to stay its own item through the merge step
    # below, rather than being folded into a neighboring cell's text.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(f"\u2022 {cell.text}")

    return lines


def _clean_lines(raw_lines: list[str]) -> list[str]:
    """Strip font-glyph artifacts, normalize whitespace, and drop noise."""
    cleaned: list[str] = []
    for line in raw_lines:
        no_artifacts = _CID_ARTIFACT_RE.sub(" ", line)
        stripped = re.sub(r"\s+", " ", no_artifacts).strip()
        if not stripped or _PAGE_NUMBER_RE.match(stripped):
            continue
        cleaned.append(stripped)
    return cleaned


def _normalize_for_header_match(bullet: str) -> str:
    """Normalize a bullet for comparison against known section headers."""
    return bullet.lower().strip(" :\u2022-")


def _process_resume_lines(lines: list[str]) -> tuple[list[Bullet], dict[str, list[Bullet]]]:
    """Merge wrapped lines, classify headings, and route bullets into sections.

    All three of these happen in a single pass, since they depend on the
    same sequential state (the current section, and whether the previous
    line was itself a section header):

    1. Reflows a bullet that wraps across two physical lines back into
       one, unless the new line is a bullet-marker-prefixed item, a
       recognized section header, or the previous line already ended in
       terminal punctuation.
    2. Classifies each freshly-started (non-continuation) bullet as a
       heading if it ends in a bare colon, or if it has no bullet marker
       *and* it's in a section where entries conventionally alternate an
       unmarked title line with marker-prefixed detail bullets (see
       _MULTI_ENTRY_SECTIONS) — this is how job-title lines, project and
       certification titles, and the contact-info block get separated
       from genuine achievement bullets. Summary, Education, and Skills
       are excluded from the marker-absence rule since they're
       conventionally flat prose or lists with no title substructure —
       without this exclusion, an unmarked Summary paragraph in a resume
       that uses markers elsewhere would be wrongly flagged as a
       heading.

       Known limitation: if a resume uses NO bullet markers anywhere at
       all, the marker signal is meaningless for it, so heading
       detection falls back to the colon-only rule (see
       resume_uses_markers below) rather than mislabeling every bullet
       as a heading.
    3. Routes each finished bullet into the current section, dropping
       the section-header line itself (e.g. "Skills") since it carries
       no content of its own beyond switching the section.
    """
    resume_uses_markers = any(_LEADING_BULLET_RE.match(line) for line in lines)

    bullets: list[Bullet] = []
    sections: dict[str, list[Bullet]] = {"content": []}
    current_section = "content"
    # Seeded True so the very first content in the document (typically a
    # name/contact-info block) is treated the same as anything else that
    # immediately follows a section switch.
    previous_was_header = True

    for line in lines:
        normalized = _normalize_for_header_match(line)
        if normalized in _HEADER_LOOKUP:
            current_section = _HEADER_LOOKUP[normalized]
            sections.setdefault(current_section, [])
            previous_was_header = True
            continue

        has_marker = bool(_LEADING_BULLET_RE.match(line))
        stripped_line = _LEADING_BULLET_RE.sub("", line).strip()
        starts_new = (
            previous_was_header
            or not bullets
            or has_marker
            or bullets[-1].text[-1:] in _TERMINAL_PUNCTUATION
        )

        if starts_new:
            is_heading = stripped_line.endswith(":") or (
                resume_uses_markers and not has_marker and current_section in _MULTI_ENTRY_SECTIONS
            )
            new_bullet = Bullet(text=stripped_line, is_heading=is_heading)
            bullets.append(new_bullet)
            sections[current_section].append(new_bullet)
        else:
            merged = Bullet(
                text=f"{bullets[-1].text} {stripped_line}".strip(),
                is_heading=bullets[-1].is_heading,
            )
            bullets[-1] = merged
            sections[current_section][-1] = merged

        previous_was_header = False

    if not sections["content"] and len(sections) > 1:
        del sections["content"]

    return bullets, sections


def _process_jd_lines(lines: list[str]) -> list[Bullet]:
    """Merge wrapped lines and classify headings for a job description.

    Job descriptions have no section structure, and — unlike resumes —
    are typically pasted as plain text with no bullet markers at all,
    even for genuine requirement sentences. So marker presence can't be
    used as a heading signal here the way it is for resumes; the only
    heading signal is a bullet ending in a bare colon (e.g. "What You'll
    Do:"), which reliably indicates an internal label rather than an
    actual requirement.
    """
    bullets: list[Bullet] = []

    for line in lines:
        has_marker = bool(_LEADING_BULLET_RE.match(line))
        stripped_line = _LEADING_BULLET_RE.sub("", line).strip()
        starts_new = not bullets or has_marker or bullets[-1].text[-1:] in _TERMINAL_PUNCTUATION

        if starts_new:
            bullets.append(Bullet(text=stripped_line, is_heading=stripped_line.endswith(":")))
        else:
            bullets[-1] = Bullet(
                text=f"{bullets[-1].text} {stripped_line}".strip(),
                is_heading=bullets[-1].is_heading,
            )

    return bullets
